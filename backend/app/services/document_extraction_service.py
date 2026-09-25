from pathlib import Path

import pandas as pd
from docx import Document
from pptx import Presentation
from pypdf import PdfReader


class DocumentExtractionService:
    """
    Extracts readable text from supported NEXUS
    research and knowledge files.

    Supported formats:
        PDF
        DOCX
        PPTX
        TXT
        CSV
    """

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".pptx",
        ".txt",
        ".csv",
    }

    def extract_text(
        self,
        file_path: str,
    ) -> str:
        """
        Extract text from a file based on its extension.
        """

        if not file_path:
            raise ValueError(
                "File path cannot be empty"
            )

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        extension = path.suffix.lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {extension}"
            )

        if extension == ".pdf":
            return self._extract_pdf(path)

        if extension == ".docx":
            return self._extract_docx(path)

        if extension == ".pptx":
            return self._extract_pptx(path)

        if extension == ".txt":
            return self._extract_txt(path)

        if extension == ".csv":
            return self._extract_csv(path)

        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    def _extract_pdf(
        self,
        path: Path,
    ) -> str:
        """
        Extract text from all pages of a PDF.
        """

        reader = PdfReader(str(path))

        pages = []

        for page in reader.pages:
            text = page.extract_text()

            if text:
                pages.append(text.strip())

        return "\n\n".join(pages).strip()

    def _extract_docx(
        self,
        path: Path,
    ) -> str:
        """
        Extract text from paragraphs in a DOCX file.
        """

        document = Document(str(path))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        return "\n\n".join(paragraphs).strip()

    def _extract_pptx(
        self,
        path: Path,
    ) -> str:
        """
        Extract text from all slides in a PPTX file.
        """

        presentation = Presentation(str(path))

        slides_text = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            slide_parts = [
                f"Slide {slide_number}:"
            ]

            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue

                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

            if len(slide_parts) > 1:
                slides_text.append(
                    "\n".join(slide_parts)
                )

        return "\n\n".join(slides_text).strip()

    def _extract_txt(
        self,
        path: Path,
    ) -> str:
        """
        Extract text from a plain-text file.
        """

        return path.read_text(
            encoding="utf-8",
            errors="replace",
        ).strip()

    def _extract_csv(
        self,
        path: Path,
    ) -> str:
        """
        Convert CSV data into readable text.
        """

        dataframe = pd.read_csv(path)

        if dataframe.empty:
            return ""

        return dataframe.to_string(
            index=False
        ).strip()


document_extraction_service = (
    DocumentExtractionService()
)